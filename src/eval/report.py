"""Отчёт сверки «ожидаемое / извлечённое» — читается жюри без нас.

    python -m src.eval.report --pred out --gold data/courts_anonymized/labels

Пишет <out>/report.md и <out>/diff.xlsx (промахи — красной заливкой, ожидаемое — в примечании ячейки).
Если рядом лежат doc_*.json, в таблицу попадают reason и evidence. Если есть routing.csv — сводка маршрутов.
Случаи из discrepancies.md класса «а» (разметка расходится с документом) помечаются отдельно и
считаются в отдельной строке, а не прячутся.
"""

from __future__ import annotations

import argparse
import csv
import json
import re
from collections import Counter
from datetime import datetime
from pathlib import Path

from src.eval.normalize import KEY_COLUMNS
from src.eval.score import FILE_COL, TABLES, read_gold, read_pred, score, summary_rows, SUMMARY_COLS
from src.eval.splits import stem

DISCREPANCY_FILES = (Path("work/discrepancies.md"), Path("docs/discrepancies.md"))
OUTCOME_RU = {"miss": "не извлечено", "fp": "лишнее", "wrong": "неверно", "ok": "✓"}


def load_doc_meta(pred: Path) -> dict[tuple[str, str], dict]:
    """(stem, столбец) → {reason, evidence, method, confidence} из doc_*.json."""
    meta = {}
    base = pred if pred.is_dir() else pred.parent
    for p in base.rglob("doc_*.json"):
        try:
            d = json.loads(p.read_text(encoding="utf-8"))
        except ValueError:
            continue
        s = stem(d.get("file") or d.get("doc_id", ""))
        for col, fv in (d.get("fields") or {}).items():
            if isinstance(fv, dict):
                meta[(s, col)] = fv
    return meta


def load_discrepancies(paths=DISCREPANCY_FILES) -> dict[tuple[str, str], dict]:
    """Строки таблицы из discrepancies.md: | файл | поле | в разметке | в документе | класс | решение |."""
    res = {}
    for p in paths:
        if not p.exists():
            continue
        for line in p.read_text(encoding="utf-8").splitlines():
            cells = [c.strip() for c in line.strip().strip("|").split("|")]
            if len(cells) < 5 or not re.search(r"(ocr|fssp)_\d+", cells[0]):
                continue
            res[(stem(cells[0]), cells[1])] = {"gold": cells[2], "doc": cells[3], "class": cells[4].lower(),
                                               "decision": cells[5] if len(cells) > 5 else ""}
    return res


def routing_summary(pred: Path) -> list[str]:
    p = (pred if pred.is_dir() else pred.parent) / "routing.csv"
    if not p.exists():
        return []
    with open(p, encoding="utf-8-sig", newline="") as f:
        rows = list(csv.DictReader(f))
    lv = Counter(r.get("level") or "—" for r in rows)
    review = sum(bool(r.get("review_flags")) for r in rows)
    return ["## Маршрутизация", "",
            f"Документов: {len(rows)} · L1: {lv.get('L1', 0)} · L2: {lv.get('L2', 0)} · L3: {lv.get('L3', 0)} · "
            f"на ручную проверку: {review}. Подробно — `routing.csv`, реестр — `registry.csv`.", ""]


def _cut(s, n=60) -> str:
    s = str(s or "").replace("|", "\\|").replace("\n", " ")
    return s if len(s) <= n else s[: n - 1] + "…"


def build_markdown(res: dict, meta: dict, disc: dict, pred: Path) -> str:
    cells = res["cells"]
    errors = [c for c in cells if not c["ok_strict"]]
    known = [c for c in errors if disc.get((c["stem"], c["column"]), {}).get("class", "").startswith("а")]
    head = [h for h, _ in SUMMARY_COLS]
    L = ["# Отчёт сверки: ожидаемое vs извлечённое", "",
         f"Сформирован {datetime.now():%Y-%m-%d %H:%M}. Предсказание `{res['pred']}`, эталон `{res['gold']}`. "
         "Сравнение после нормализации (деньги — 2 знака, даты — ГГГГ-ММ-ДД, пусто == пусто, 0 ≠ пусто).", "",
         "## Итог по срезам", "", "| " + " | ".join(head) + " |", "|" + "---|" * len(head)]
    L += ["| " + " | ".join(r) + " |" for r in summary_rows(res)]
    L += [""] + routing_summary(pred)

    L += ["## Где ошибаемся чаще всего", "",
          "| таблица | столбец | ошибок | не извлечено | лишнее | неверно | из них расхождение эталона |",
          "|---|---|---|---|---|---|---|"]
    by_col = Counter((c["table"], c["column"]) for c in errors)
    for (t, col), n in by_col.most_common():
        cs = [c for c in errors if (c["table"], c["column"]) == (t, col)]
        kn = sum(c in known for c in cs)
        name = f"**{col}**" if col in KEY_COLUMNS else col
        L.append(f"| {t} | {name} | {n} | {sum(c['outcome'] == 'miss' for c in cs)} | "
                 f"{sum(c['outcome'] == 'fp' for c in cs)} | {sum(c['outcome'] == 'wrong' for c in cs)} | {kn} |")
    if not by_col:
        L.append("| — | расхождений нет | 0 | 0 | 0 | 0 | 0 |")
    L += ["", f"Всего ошибочных ячеек: {len(errors)} из {len(cells)}; из них {len(known)} — "
          "разметка расходится с текстом документа (класс «а» в discrepancies.md, значение оставлено нашим).", ""]

    L += ["## По документам", ""]
    clean = []
    for s in sorted({c["stem"] for c in cells}):
        ds = [c for c in cells if c["stem"] == s]
        bad = [c for c in ds if not c["ok_strict"]]
        if not bad:
            clean.append(s)
            continue
        f = ds[0]["file"]
        absent = "" if ds[0]["present"] else " — **документ не обработан**"
        L += [f"### {s} · `{f}` · ошибок {len(bad)} из {len(ds)}{absent}", "",
              "| столбец | ожидалось | извлечено | итог | причина | цитата |", "|---|---|---|---|---|---|"]
        for c in bad:
            m = meta.get((s, c["column"]), {})
            d = disc.get((s, c["column"]))
            verdict = OUTCOME_RU[c["outcome"]] + (" (loose ✓)" if c["ok_loose"] else "")
            if d:
                verdict += f" · эталон: класс «{d['class']}»"
            L.append(f"| {c['column']} | {_cut(c['gold'])} | {_cut(c['pred'])} | {verdict} | "
                     f"{_cut(m.get('reason'), 30)} | {_cut(m.get('evidence'), 70)} |")
        L.append("")
    if clean:
        L += [f"Без единой ошибки ({len(clean)}): " + ", ".join(clean), ""]
    return "\n".join(L)


def write_diff_xlsx(res: dict, pred_rows: dict, gold_rows: dict, path: Path) -> None:
    from openpyxl import Workbook                                  # pylint: disable=import-outside-toplevel
    from openpyxl.comments import Comment
    from openpyxl.styles import Font, PatternFill

    red = PatternFill("solid", start_color="F8CBAD")
    grey = PatternFill("solid", start_color="EDEDED")
    bold = Font(name="Arial", bold=True)
    normal = Font(name="Arial")
    status = {(c["stem"], c["column"]): c for c in res["cells"]}
    wb = Workbook()
    wb.remove(wb.active)
    for t in TABLES:
        if not gold_rows.get(t):
            continue
        ws = wb.create_sheet(t)
        cols = list(gold_rows[t][0].keys())
        ws.append(cols)
        for cell in ws[1]:
            cell.font = bold
        pred_by = {stem(r.get(FILE_COL[t], "")): r for r in pred_rows.get(t, [])}
        for g in gold_rows[t]:
            s = stem(g[FILE_COL[t]])
            p = pred_by.get(s)
            ws.append([g[FILE_COL[t]] if col == FILE_COL[t] else ("" if p is None else p.get(col, ""))
                       for col in cols])
            r = ws.max_row
            for j, col in enumerate(cols, 1):
                cell = ws.cell(r, j)
                cell.font = normal
                c = status.get((s, col))
                if p is None:
                    cell.fill = grey
                elif c and not c["ok_strict"]:
                    cell.fill = red
                    cell.comment = Comment(f"ожидалось: {c['gold'] or '(пусто)'}", "scorer")
        ws.freeze_panes = "B2"
        for j, col in enumerate(cols, 1):
            ws.column_dimensions[ws.cell(1, j).column_letter].width = 40 if col in ("IdDebtText", "DbtrAdr") else 18
    legend = wb.create_sheet("легенда")
    for row in (["красная заливка", "значение не совпало с эталоном; ожидаемое — в примечании ячейки"],
                ["серая заливка", "документ отсутствует в предсказании"],
                ["сравнение", "после нормализации: деньги 2 знака, даты ГГГГ-ММ-ДД, пусто == пусто"]):
        legend.append(row)
    wb.save(path)


def main() -> None:
    ap = argparse.ArgumentParser()
    ap.add_argument("--pred", default="out")
    ap.add_argument("--gold", default="data/courts_anonymized/labels")
    ap.add_argument("-o", "--out", default=None, help="папка для report.md / diff.xlsx; по умолчанию — --pred")
    ap.add_argument("--missing-as-error", action="store_true")
    a = ap.parse_args()
    pred, gold = Path(a.pred), Path(a.gold)
    out = Path(a.out) if a.out else (pred if pred.is_dir() else pred.parent)
    out.mkdir(parents=True, exist_ok=True)
    res = score(pred, gold, missing_as_error=a.missing_as_error)
    md = build_markdown(res, load_doc_meta(pred), load_discrepancies(), pred)
    (out / "report.md").write_text(md, encoding="utf-8")
    write_diff_xlsx(res, read_pred(pred), read_gold(gold), out / "diff.xlsx")
    print(f"→ {out / 'report.md'}, {out / 'diff.xlsx'}")


if __name__ == "__main__":
    main()
