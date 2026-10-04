"""Чистые функции для UI: без streamlit, чтобы их можно было проверить тестами и переиспользовать.

Doc берётся через duck typing (pydantic-Doc из src.core.contract или dict из doc_*.json)."""

from __future__ import annotations

import csv
import io
import json
import os
import urllib.request
from pathlib import Path

ROOT = Path(__file__).resolve().parents[2]

LEVEL_INFO = {
    "L1": ("#b42318", "#fdecea", "Срочно", "есть срок не больше 10 дней (для судебных актов — 15) — письмо уходит сразу"),
    "L2": ("#a15c07", "#fef3e2", "Нужно действие", "юристу нужно действовать, но срок не горит"),
    "L3": ("#475467", "#f2f4f7", "В реестр", "действий не требуется, только строка в реестре"),
}

REVIEW_TEXT_FALLBACK = {
    "LOW_CONFIDENCE": "низкая уверенность распознавания",
    "REQUIRED_FIELD_MISSING": "не извлечены обязательные реквизиты",
    "ID_CONFLICT": "противоречия в реквизитах (несколько ИНН/паспортов)",
    "ARITH_CHECK_FAILED": "сумма не сходится с составляющими",
    "UNCLASSIFIED": "тип документа не распознан",
}

REASON_TEXT = {
    "NOT_IN_TEXT": "в документе нет",
    "ANCHOR_NOT_FOUND": "якорь не найден",
    "OCR_UNREADABLE": "не читается после OCR",
    "VALIDATOR_FAILED": "не прошло проверку",
    "LLM_DISAGREE": "LLM не согласна с правилом",
    "LLM_DISABLED": "нужна LLM, она выключена",
    "LLM_FALLBACK_CANDIDATE": "кандидат на дозаполнение LLM",
}

METHOD_TEXT = {"copy": "копия тега XML", "lookup": "справочник", "regex": "правило", "anchor": "якорь",
               "llm": "LLM", "derived": "вычислено", "gold": "эталон (фикстура)", "none": "—"}

SUPPORTED = (".pdf", ".xml", ".json")


def get(obj, key, default=None):
    if obj is None:
        return default
    if isinstance(obj, dict):
        return obj.get(key, default)
    return getattr(obj, key, default)


# ---------------------------------------------------------------- расшифровки кодов
def code_texts() -> dict[str, str]:
    """Код причины → что сделать юристу: rules/codes.LETTER_TEXT (junior ML) + действия из маркеров актов."""
    out: dict[str, str] = {}
    try:
        from src.rules.markers import MARKERS  # pylint: disable=import-outside-toplevel
        out.update({code: action for _, code, _, action, _ in MARKERS})
    except ImportError:
        pass
    try:
        from src.rules.codes import LETTER_TEXT  # pylint: disable=import-outside-toplevel
        out.update(LETTER_TEXT)
    except ImportError:
        pass
    return out


def review_texts() -> dict[str, str]:
    try:
        from src.export.mailer import REVIEW_TEXT  # pylint: disable=import-outside-toplevel
        return {**REVIEW_TEXT_FALLBACK, **REVIEW_TEXT}
    except ImportError:
        return REVIEW_TEXT_FALLBACK


# ---------------------------------------------------------------- LLM
def ping_models(base_url: str, timeout: float = 2.0) -> tuple[bool, list[str], str | None]:
    """GET {base_url}/models (Ollama, OpenAI-совместимый). (доступна, [модели], ошибка)."""
    url = base_url.rstrip("/") + "/models"
    try:
        with urllib.request.urlopen(url, timeout=timeout) as r:            # noqa: S310 — локальный адрес из .env
            data = json.loads(r.read().decode("utf-8"))
        models = sorted(m.get("id", "") for m in data.get("data", []) if m.get("id"))
        return True, models, None
    except Exception as e:                                                  # pylint: disable=broad-except
        return False, [], f"{type(e).__name__}"


def set_llm(enabled: bool) -> None:
    """Переключатель «без LLM» в UI. В клиенте есть только disable(), поэтому выставляем флаг напрямую."""
    from src.llm import client  # pylint: disable=import-outside-toplevel
    if hasattr(client, "enable") and enabled:
        client.enable()
    elif not enabled:
        client.disable()
    else:
        client._DISABLED = False  # pylint: disable=protected-access
    os.environ["NO_LLM"] = "0" if enabled else "1"


# ---------------------------------------------------------------- документы
def list_examples(root: Path = ROOT) -> dict[str, list[Path]]:
    groups = {
        "Постановления ФССП (XML)": sorted((root / "data/courts_anonymized/fssp").rglob("*.xml")),
        "Приказы и ИЛ (PDF)": sorted((root / "data/courts_anonymized/ocr").rglob("*.pdf")),
        "Судебные акты (PDF)": sorted((root / "data/acts").rglob("*.pdf")),
        "Эталонные doc.json (fixtures)": sorted((root / "fixtures").glob("doc_*.json")),
    }
    return {k: v for k, v in groups.items() if v}


def precomputed_docs(out_dir: Path) -> list[Path]:
    """doc_*.json, которые run_examples сохранил в out/ (предрасчитанные результаты для режима без LLM)."""
    return sorted(Path(out_dir).rglob("doc_*.json")) if Path(out_dir).exists() else []


def fallback_json(stem: str, out_dir: Path, root: Path = ROOT) -> Path | None:
    """Если ветка ещё не подключена — ищем готовый результат по stem: сначала out/, потом fixtures/."""
    for p in (*precomputed_docs(out_dir), *sorted((root / "fixtures").glob("doc_*.json"))):
        if p.stem.removeprefix("doc_") == stem:
            return p
    return None


def field_rows(doc, key_columns: set[str] | None = None) -> list[dict]:
    key_columns = key_columns or set()
    rows = []
    for name, fv in (get(doc, "fields") or {}).items():
        value = get(fv, "value") or ""
        conf = get(fv, "confidence")
        reason = get(fv, "reason")
        rows.append({
            "поле": name,
            "значение": value,
            "уверенность": None if conf is None else round(float(conf), 2),
            "метод": METHOD_TEXT.get(get(fv, "method") or "none", get(fv, "method")),
            "причина пустого": "" if value else REASON_TEXT.get(reason, reason or ""),
            "цитата": get(fv, "evidence") or "",
            "_key": name in key_columns,
            "_llm": get(fv, "method") == "llm",
        })
    return rows


def fill_stats(doc) -> tuple[int, int]:
    fields = get(doc, "fields") or {}
    filled = sum(1 for fv in fields.values() if get(fv, "value"))
    return filled, len(fields)


def table_columns(doc) -> list[str]:
    from src.core.columns import ACT_FIELDS, OCR_COLUMNS, XML_COLUMNS  # pylint: disable=import-outside-toplevel
    return {"ocr": OCR_COLUMNS, "xml": XML_COLUMNS, "acts": ACT_FIELDS}[get(doc, "table")]


def table_row(doc, cols: list[str] | None = None) -> dict:
    """Строка в формате шаблона оргов (ocr.csv / xml.csv) + маршрут."""
    cols = cols or table_columns(doc)
    fields = get(doc, "fields") or {}
    row = {c: get(fields.get(c), "value") or "" for c in cols}
    t = get(doc, "table")
    if t == "ocr":
        row["Файл"], row["тип документа"] = get(doc, "file"), get(doc, "doc_type")
    elif t == "xml":
        row["FileName"] = get(doc, "file")
    route = get(doc, "route")
    lawyer, review = get(route, "lawyer"), get(route, "review")
    row["route_level"] = get(lawyer, "level") or ""
    row["route_codes"] = ";".join(get(lawyer, "reason_codes") or [])
    row["review_flags"] = ";".join(get(review, "flags") or [])
    return row


def rows_to_csv(rows: list[dict]) -> bytes:
    """UTF-8 с BOM, запятая, CRLF — как шаблоны оргов."""
    if not rows:
        return b""
    fields = list(dict.fromkeys(k for r in rows for k in r))
    buf = io.StringIO()
    w = csv.DictWriter(buf, fieldnames=fields, lineterminator="\r\n")
    w.writeheader()
    w.writerows(rows)
    return ("\ufeff" + buf.getvalue()).encode("utf-8")


def rows_to_xlsx(rows: list[dict], sheet: str = "результат") -> bytes:
    """xlsx: всё как текст (номера, паспорта, ИНН сохраняют ведущие нули)."""
    from openpyxl import Workbook  # pylint: disable=import-outside-toplevel
    from openpyxl.styles import Font
    wb = Workbook()
    ws = wb.active
    ws.title = sheet[:31]
    if rows:
        fields = list(dict.fromkeys(k for r in rows for k in r))
        ws.append(fields)
        for c in ws[1]:
            c.font = Font(name="Arial", bold=True)
        for r in rows:
            ws.append([str(r.get(f, "") or "") for f in fields])
        for row in ws.iter_rows(min_row=2):
            for c in row:
                c.number_format = "@"
                c.font = Font(name="Arial")
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def doc_json_bytes(doc) -> bytes:
    data = doc.model_dump() if hasattr(doc, "model_dump") else doc
    return json.dumps(data, ensure_ascii=False, indent=2).encode("utf-8")


# ---------------------------------------------------------------- реестр и качество
def read_csv_rows(path: Path) -> list[dict]:
    if not Path(path).exists():
        return []
    with open(path, encoding="utf-8-sig", newline="") as f:
        return list(csv.DictReader(f))


def registry_rows(out_dir: Path) -> tuple[list[dict], str]:
    """(строки, источник): registry.csv, а пока его нет — routing.csv."""
    out_dir = Path(out_dir)
    for name in ("registry.csv", "routing.csv"):
        rows = read_csv_rows(out_dir / name)
        if rows:
            for r in rows:                       # routing.csv называет столбцы иначе — приводим к реестру
                r.setdefault("route_level", r.get("level", ""))
                r.setdefault("route_codes", r.get("reason_codes", ""))
            return rows, name
    return [], ""
